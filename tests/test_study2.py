"""Tests for tools/pipeline/study2_reliability.py on synthetic workbooks.

Two small rating workbooks (pass A = Rater 1, pass B = Rater 2) are generated with
openpyxl using the template columns: three made-up recall clips for participants p01
and p02, made-up transcripts and hand scores, plus one letter-fluency row in the
Transcriptions tab that the recall-only script must ignore. Nothing here comes from a
real recording.
"""
from __future__ import annotations

import importlib
import json
import os
import subprocess
import sys

import pandas as pd
import pytest
from openpyxl import Workbook

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SCRIPT = os.path.join(REPO, "tools", "pipeline", "study2_reliability.py")
sys.path.insert(0, os.path.dirname(SCRIPT))
study2 = importlib.import_module("study2_reliability")

TRANSCRIPT_COLS = ["clip_id", "audio_file", "participant", "date", "time", "task",
                   "cue (letter/category)", "transcription", "your_initials", "notes"]
RECALL_COLS = ["clip_id", "audio_file", "participant", "date", "time", "list",
               "words shown (in the order presented)", "your transcription (filled in for you)",
               "Correct", "Repetitions", "Intrusions", "Primacy", "Middle", "Recency",
               "Serial Cluster", "notes"]

# Three synthetic recall clips; the third file name carries spaces so the join has to strip them.
CLIPS = [
    {"clip_id": "R0001", "audio_file": "clip_p01_recall_01.m4a", "participant": "p01",
     "task": "word-list recall", "text": "apple river stone cloud pencil",
     "scores": {"Correct": 5, "Repetitions": 1, "Intrusions": 0, "Primacy": 2, "Middle": 2, "Recency": 1}},
    {"clip_id": "R0002", "audio_file": "clip_p01_recall_02.m4a", "participant": "p01",
     "task": "word-list recall", "text": "candle mirror garden window ladder rope",
     "scores": {"Correct": 8, "Repetitions": 0, "Intrusions": 2, "Primacy": 3, "Middle": 3, "Recency": 2}},
    {"clip_id": "R0003", "audio_file": "clip p02 recall 01.m4a", "participant": "p02",
     "task": "word-list recall", "text": "basket tunnel um feather",
     "scores": {"Correct": 3, "Repetitions": 2, "Intrusions": 1, "Primacy": 1, "Middle": 1, "Recency": 1}},
]
OTHER_TASK = {"clip_id": "R0004", "audio_file": "clip_p02_letter_01.m4a", "participant": "p02",
              "task": "letter fluency", "text": "bottle bridge button"}
MEASURES = ["recall_correct", "recall_repetitions", "recall_intrusions",
            "recall_primacy", "recall_middle", "recall_recency"]
SCORE_KEYS = ["Correct", "Repetitions", "Intrusions", "Primacy", "Middle", "Recency"]


def basename(clip: dict) -> str:
    return clip["audio_file"].replace(" ", "")


def build_workbook(path: str, pass_label: str, overrides: dict | None = None) -> None:
    """overrides maps (clip_id, column) to a replacement value, or (clip_id, 'transcription');
    None leaves the cell blank."""
    overrides = overrides or {}

    def val(clip, col, default):
        return overrides.get((clip["clip_id"], col), default)

    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    ws.cell(row=1, column=1, value=f"Synthetic reliability set - PASS {pass_label}")
    ws = wb.create_sheet("Transcriptions")
    ws.append(TRANSCRIPT_COLS)
    for c in CLIPS + [OTHER_TASK]:
        ws.append([c["clip_id"], c["audio_file"], c["participant"], "2030-01-01", "09:00:00",
                   c["task"], "", val(c, "transcription", c["text"]), "xx", ""])
    ws = wb.create_sheet("Score - word recall")
    ws.append(RECALL_COLS)
    for i, c in enumerate(CLIPS):
        row = [c["clip_id"], c["audio_file"], c["participant"], "2030-01-01", "09:00:00",
               "list 1", "1. apple, 2. river", f'=IFERROR(VLOOKUP($A{i + 2},Transcriptions!$A:$J,8,FALSE),"")']
        row += [val(c, k, c["scores"][k]) for k in SCORE_KEYS] + ["", ""]
        ws.append(row)
    wb.save(path)


def build_machine(root: str, skip_transcripts=()) -> dict:
    """Machine tables that reproduce pass A exactly, plus one transcript JSON per clip."""
    os.makedirs(os.path.join(root, "transcripts"), exist_ok=True)
    recall_rows, serial_rows = [], []
    for c in CLIPS:
        b = basename(c)
        if c["clip_id"] not in skip_transcripts:
            with open(os.path.join(root, "transcripts", b + ".json"), "w") as fh:
                json.dump({"basename": b, "text": c["text"], "model": "synthetic"}, fh)
        s = c["scores"]
        recall_rows.append({"n_target": 12, "n_correct": s["Correct"], "n_repetitions": s["Repetitions"],
                            "n_intrusions": s["Intrusions"], "basename": b,
                            "recipe_version": "2026-09-freeze", "recipe_config": "ra_informed"})
        serial_rows.append({"basename": b, "primacy_3_6_3": s["Primacy"], "middle_3_6_3": s["Middle"],
                            "recency_3_6_3": s["Recency"]})
    paths = {"recall": os.path.join(root, "scores_word_recall.csv"),
             "serial": os.path.join(root, "scores_serial.csv"),
             "transcripts": os.path.join(root, "transcripts")}
    pd.DataFrame(recall_rows).to_csv(paths["recall"], index=False)
    pd.DataFrame(serial_rows).to_csv(paths["serial"], index=False)
    return paths


def run(tmp_path, overrides_b=None, dry_run=False, **machine_kw):
    a, b = str(tmp_path / "passA.xlsx"), str(tmp_path / "passB.xlsx")
    build_workbook(a, "A")
    build_workbook(b, "B", overrides_b)
    paths = build_machine(str(tmp_path / "machine"), **machine_kw)
    out = str(tmp_path / "study2")
    argv = ["--pass-a", a, "--pass-b", b, "--scores-recall", paths["recall"],
            "--transcripts-dir", paths["transcripts"], "--out-dir", out]
    if dry_run:
        argv.append("--dry-run")
    return study2.main(argv), out, argv


def load_json(path: str) -> dict:
    with open(path) as fh:
        return json.load(fh)


def stat(table: pd.DataFrame, measure: str, comparison: str, col: str):
    sub = table[(table["measure"] == measure) & (table["comparison"] == comparison)]
    assert len(sub) == 1, (measure, comparison)
    return sub[col].iloc[0]


def wer_stat(table: pd.DataFrame, comparison: str, col: str):
    sub = table[(table["task"] == "word_recall") & (table["comparison"] == comparison)]
    assert len(sub) == 1, comparison
    return sub[col].iloc[0]


def test_identical_passes_give_perfect_agreement(tmp_path):
    code, out, _ = run(tmp_path)
    assert code == 0
    scores = pd.read_csv(os.path.join(out, "study2_score_agreement.csv"))
    wer = pd.read_csv(os.path.join(out, "study2_transcript_agreement.csv"))
    assert len(scores) == 6 * 4
    assert set(scores["measure"]) == set(MEASURES) == {m[0] for m in study2.MEASURES}
    assert set(scores["comparison"]) == set(study2.COMPARISONS)
    assert len(wer) == 4 and set(wer["task"]) == {"word_recall"}
    for m in MEASURES:
        for comp in study2.COMPARISONS:
            assert stat(scores, m, comp, "icc21") == 1.0, (m, comp)
            assert stat(scores, m, comp, "mae") == 0.0, (m, comp)
            assert stat(scores, m, comp, "bias") == 0.0, (m, comp)
        assert stat(scores, m, "passA_vs_passB", "n") == 3
    assert (wer["mean_wer"].fillna(0) == 0).all()
    assert wer_stat(wer, "passA_vs_passB", "n") == 3     # the letter-fluency row is ignored
    for name in ["study2_report.md", "study2_results.json", "study2_agreement_by_measure.png",
                 "study2_paired_scores.csv", "study2_paired_transcripts.csv"]:
        assert os.path.exists(os.path.join(out, name)), name
    with open(os.path.join(out, "study2_report.md")) as fh:
        report = fh.read()
    assert "clip_p01" not in report and "clipp02" not in report  # aggregate only
    assert "—" not in report
    results = load_json(os.path.join(out, "study2_results.json"))
    assert len(results["score_agreement"]) == 24
    assert results["recipe"] == {"recipe_version": ["2026-09-freeze"], "recipe_config": ["ra_informed"]}
    assert results["matching"]["recall_correct"]["with_machine_row"] == 3
    assert results["data_quality"]["machine_transcripts_missing"] == 0
    assert all(os.sep not in v for v in results["inputs"].values())  # file names only, no paths
    assert os.path.getsize(os.path.join(out, "study2_agreement_by_measure.png")) > 1000


def test_known_differences_and_bad_cells(tmp_path):
    overrides = {
        ("R0001", "Correct"): 6,             # +1
        ("R0003", "Correct"): 5,             # +2  -> MAE 1.0, bias +1.0 over 3 clips
        ("R0002", "Repetitions"): "two",     # non-numeric, dropped
        ("R0003", "Intrusions"): None,       # blank, dropped
        ("R0001", "transcription"): "apple river stone cloud pen",  # 1 of 5 words differs
    }
    code, out, _ = run(tmp_path, overrides_b=overrides, skip_transcripts=("R0002",))
    assert code == 0
    scores = pd.read_csv(os.path.join(out, "study2_score_agreement.csv"))
    wer = pd.read_csv(os.path.join(out, "study2_transcript_agreement.csv"))
    results = load_json(os.path.join(out, "study2_results.json"))
    assert stat(scores, "recall_correct", "passA_vs_passB", "n") == 3
    assert stat(scores, "recall_correct", "passA_vs_passB", "mae") == pytest.approx(1.0)
    assert stat(scores, "recall_correct", "passA_vs_passB", "bias") == pytest.approx(1.0)
    assert stat(scores, "recall_correct", "machine_vs_passA", "mae") == 0.0
    assert stat(scores, "recall_correct", "machine_vs_passB", "bias") == pytest.approx(-1.0)
    assert stat(scores, "recall_correct", "machine_vs_human_mean", "bias") == pytest.approx(-0.5)
    assert stat(scores, "recall_repetitions", "passA_vs_passB", "n") == 2
    assert stat(scores, "recall_intrusions", "passA_vs_passB", "n") == 2
    assert stat(scores, "recall_primacy", "passA_vs_passB", "n") == 3
    qb = results["data_quality"]["passB"]
    assert qb["nonnumeric"]["Score - word recall / Repetitions"] == 1
    assert qb["blank"]["Score - word recall / Intrusions"] == 1
    # transcripts: one substituted word in five on one of three clips; one machine transcript missing
    assert wer_stat(wer, "passA_vs_passB", "mean_wer") == pytest.approx(0.2 / 3)
    assert wer_stat(wer, "passA_vs_passB", "corpus_wer") == pytest.approx(1 / 15)
    assert wer_stat(wer, "machine_vs_passA", "n") == 2
    assert wer_stat(wer, "machine_vs_passA", "mean_wer") == 0.0
    assert wer_stat(wer, "machine_vs_passB", "mean_wer") == pytest.approx(0.1)
    assert results["data_quality"]["machine_transcripts_missing"] == 1


def test_dry_run_validates_and_writes_nothing(tmp_path):
    code, out, argv = run(tmp_path, dry_run=True)
    assert code == 0 and not os.path.exists(out)
    proc = subprocess.run([sys.executable, SCRIPT] + argv, capture_output=True, text=True, cwd=REPO, check=False)
    assert proc.returncode == 0, proc.stderr
    assert "dry run" in proc.stdout and "Recall correct" in proc.stdout and "n=3" in proc.stdout
    assert not os.path.exists(out)


def test_missing_column_is_reported_not_raised(tmp_path, capsys):
    a, b = str(tmp_path / "passA.xlsx"), str(tmp_path / "passB.xlsx")
    build_workbook(a, "A")
    build_workbook(b, "B")
    from openpyxl import load_workbook
    wb = load_workbook(b)
    wb["Score - word recall"].cell(row=1, column=RECALL_COLS.index("Correct") + 1, value="Right")
    wb.save(b)
    paths = build_machine(str(tmp_path / "machine"))
    code = study2.main(["--pass-a", a, "--pass-b", b, "--scores-recall", paths["recall"],
                        "--transcripts-dir", paths["transcripts"], "--out-dir", str(tmp_path / "study2")])
    assert code == 2
    assert "lacks column 'Correct'" in capsys.readouterr().out
    assert not os.path.exists(tmp_path / "study2")


def test_untouched_templates_run_with_zero_n(tmp_path):
    """The templates as sent: every transcription and score cell blank."""
    blank = {}
    for c in CLIPS + [OTHER_TASK]:
        blank[(c["clip_id"], "transcription")] = None
    for c in CLIPS:
        for k in c["scores"]:
            blank[(c["clip_id"], k)] = None
    a, b = str(tmp_path / "passA.xlsx"), str(tmp_path / "passB.xlsx")
    build_workbook(a, "A", blank)
    build_workbook(b, "B", blank)
    paths = build_machine(str(tmp_path / "machine"))
    out = str(tmp_path / "study2")
    code = study2.main(["--pass-a", a, "--pass-b", b, "--scores-recall", paths["recall"],
                        "--transcripts-dir", paths["transcripts"], "--out-dir", out])
    assert code == 0
    scores = pd.read_csv(os.path.join(out, "study2_score_agreement.csv"))
    wer = pd.read_csv(os.path.join(out, "study2_transcript_agreement.csv"))
    assert len(scores) == 24 and (scores["n"] == 0).all()
    assert len(wer) == 4 and (wer["n"] == 0).all()
    results = load_json(os.path.join(out, "study2_results.json"))
    assert results["data_quality"]["passA"]["unscored_rows"]["Score - word recall"] == 3
    assert results["data_quality"]["passA"]["unscored_rows"]["Transcriptions"] == 3


def test_helpers():
    assert study2.clean_basename(" clip p02 recall 01.m4a ") == "clipp02recall01.m4a"
    assert study2.clean_basename(None) == "" and study2.clean_basename(float("nan")) == ""
    assert study2.canon_task("Word-List Recall") == "word_recall"
    assert study2.canon_task("") is None
    vals, n_blank, n_nonnum = study2.coerce_numeric(pd.Series([3, "4", " 5 ", "", None, "x", "1 2"]))
    assert vals.tolist()[:3] == [3.0, 4.0, 5.0]
    assert (n_blank, n_nonnum) == (2, 2)
